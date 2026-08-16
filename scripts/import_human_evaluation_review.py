from __future__ import annotations

# ruff: noqa: I001

import argparse
import asyncio
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
BACKEND_ROOT = REPOSITORY_ROOT / "backend"
sys.path.insert(0, str(BACKEND_ROOT))

from app.config import Settings  # noqa: E402
from app.database.session import DatabaseManager  # noqa: E402
from app.evaluation.human_review import (  # noqa: E402
    WORKBOOK_COLUMNS,
    TraceabilityIndex,
    execute_human_review,
    parse_review_rows,
    validate_review_rows,
)
from app.models import Chunk, Document, Source  # noqa: E402
from app.repositories.evaluation import EvaluationRepository  # noqa: E402


def _arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Validate or apply the human-reviewed Phase J evaluation workbook"
    )
    parser.add_argument(
        "workbook",
        nargs="?",
        type=Path,
        default=REPOSITORY_ROOT / "data" / "evaluation" / "eval_human_review.xlsx",
    )
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument(
        "--dry-run",
        action="store_true",
        help="Validate only (default); never writes to PostgreSQL or output JSON.",
    )
    mode.add_argument(
        "--apply",
        action="store_true",
        help="Persist review decisions; final JSON is written only after every row is reviewed.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=REPOSITORY_ROOT / "data" / "evaluation" / "phase_j_human_verified.json",
    )
    return parser.parse_args()


def _workbook_records(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = (
            workbook["Human Review"] if "Human Review" in workbook.sheetnames else workbook.active
        )
        iterator = sheet.iter_rows(values_only=True)
        headers = tuple(str(value).strip() if value is not None else "" for value in next(iterator))
        missing = [column for column in WORKBOOK_COLUMNS if column not in headers]
        if missing:
            raise ValueError(f"missing workbook columns: {', '.join(missing)}")
        records: list[dict[str, object]] = []
        for values in iterator:
            record = dict(zip(headers, values, strict=False))
            if any(value not in (None, "") for value in record.values()):
                records.append(record)
        return records
    finally:
        workbook.close()


async def _run(workbook_path: Path, output_path: Path, *, apply: bool) -> dict[str, object]:
    records = _workbook_records(workbook_path)
    rows = parse_review_rows(records)
    database = DatabaseManager(Settings())
    try:
        async with database.session_factory() as session:
            document_rows = (
                await session.execute(
                    select(
                        Document.document_id,
                        Document.source_url,
                        Document.title,
                        Source.name,
                    ).join(Source, Source.id == Document.source_id)
                )
            ).all()
            chunk_rows = (
                await session.execute(
                    select(Chunk.chunk_id, Document.document_id, Chunk.content).join(
                        Document, Document.id == Chunk.document_id
                    )
                )
            ).all()
            traceability = TraceabilityIndex(
                document_urls={
                    document_id: source_url
                    for document_id, source_url, _title, _source_title in document_rows
                },
                chunk_documents={
                    chunk_id: document_id for chunk_id, document_id, _content in chunk_rows
                },
                chunk_contents={
                    chunk_id: content for chunk_id, _document_id, content in chunk_rows
                },
                document_titles={
                    document_id: title
                    for document_id, _source_url, title, _source_title in document_rows
                },
                source_titles={
                    document_id: source_title or ""
                    for document_id, _source_url, _title, source_title in document_rows
                },
            )
            validation = validate_review_rows(rows, traceability)

            async def persist(questions: tuple[Any, ...]) -> None:
                repository = EvaluationRepository(session)
                await repository.upsert_questions(questions)
                await repository.commit()

            execution = await execute_human_review(
                validation,
                apply=apply,
                output_path=output_path,
                persist=persist,
            )
            return execution.to_dict()
    finally:
        await database.dispose()


def main() -> None:
    arguments = _arguments()
    result = asyncio.run(_run(arguments.workbook, arguments.output, apply=bool(arguments.apply)))
    print(json.dumps(result, ensure_ascii=True, sort_keys=True))
    if result["TRACEABILITY_ERRORS"]:
        raise SystemExit(2)


if __name__ == "__main__":
    main()
