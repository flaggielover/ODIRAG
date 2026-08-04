from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.parsers.base import ParsedArtifact, ParsedTable, table_to_markdown


class XlsxParser:
    def __init__(self, *, max_rows_per_sheet: int = 10_000) -> None:
        self.max_rows_per_sheet = max_rows_per_sheet

    def parse(self, content: bytes, *, filename: str | None = None) -> ParsedArtifact:
        del filename
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        tables: list[ParsedTable] = []
        text_parts: list[str] = []
        sheet_metadata: list[dict[str, Any]] = []
        try:
            for worksheet in workbook.worksheets:
                rows = [
                    [_stringify(value) for value in row]
                    for row in worksheet.iter_rows(
                        min_row=1, max_row=self.max_rows_per_sheet, values_only=True
                    )
                ]
                rows = [row for row in rows if any(value for value in row)]
                if not rows:
                    sheet_metadata.append({"name": worksheet.title, "row_count": 0})
                    continue
                headers, body = rows[0], rows[1:]
                markdown = table_to_markdown(headers, body)
                table = ParsedTable(
                    tuple(headers),
                    tuple(tuple(value for value in row) for row in body),
                    markdown,
                )
                tables.append(table)
                text_parts.append(f"# {worksheet.title}\n\n{markdown}")
                sheet_metadata.append(
                    {
                        "name": worksheet.title,
                        "row_count": len(body),
                        "headers": headers,
                        "rows": [dict(zip(headers, row, strict=False)) for row in body],
                    }
                )
        finally:
            workbook.close()
        return ParsedArtifact(
            text="\n\n".join(text_parts),
            tables=tuple(tables),
            metadata={"sheets": sheet_metadata},
        )


def _stringify(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()
