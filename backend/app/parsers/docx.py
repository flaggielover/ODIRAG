from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.oxml.ns import qn

from app.parsers.base import ParsedArtifact, ParsedSection, ParsedTable, table_to_markdown


class DocxParser:
    def parse(self, content: bytes, *, filename: str | None = None) -> ParsedArtifact:
        del filename
        document = Document(BytesIO(content))
        sections: list[ParsedSection] = []
        tables: list[ParsedTable] = []
        text_parts: list[str] = []
        current_title: str | None = None
        current_level = 0
        current_parts: list[str] = []
        hyperlinks: list[dict[str, str]] = []
        for paragraph in document.paragraphs:
            value = paragraph.text.strip()
            style_name = paragraph.style.name if paragraph.style is not None else ""
            if style_name.lower().startswith("heading") and value:
                if current_title is not None or current_parts:
                    sections.append(
                        ParsedSection(
                            current_title, current_level, "\n".join(current_parts).strip()
                        )
                    )
                current_level = _heading_level(style_name)
                current_title = value
                current_parts = []
                text_parts.append("#" * current_level + " " + value)
            elif value:
                current_parts.append(value)
                text_parts.append(value)
            hyperlinks.extend(_paragraph_hyperlinks(paragraph, document.part.rels))
        if current_title is not None or current_parts:
            sections.append(
                ParsedSection(current_title, current_level, "\n".join(current_parts).strip())
            )
        for table in document.tables:
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            headers, body = rows[0], rows[1:]
            parsed = ParsedTable(
                tuple(headers),
                tuple(tuple(value for value in row) for row in body),
                table_to_markdown(headers, body),
            )
            tables.append(parsed)
            text_parts.append(parsed.markdown)
        return ParsedArtifact(
            text="\n\n".join(text_parts),
            sections=tuple(sections),
            tables=tuple(tables),
            metadata={"hyperlinks": hyperlinks},
        )


def _heading_level(style_name: str) -> int:
    digits = "".join(character for character in style_name if character.isdigit())
    return max(1, min(6, int(digits))) if digits else 1


def _paragraph_hyperlinks(paragraph: Any, relationships: Any) -> list[dict[str, str]]:
    links = []
    for hyperlink in paragraph._p.xpath(".//w:hyperlink"):
        relationship_id = hyperlink.get(qn("r:id"))
        if relationship_id and relationship_id in relationships:
            text = "".join(node.text or "" for node in hyperlink.xpath(".//w:t"))
            links.append({"text": text, "href": str(relationships[relationship_id].target_ref)})
    return links
