from __future__ import annotations

from datetime import date
from importlib import import_module
from io import BytesIO
from zipfile import ZipFile

from docx import Document as DocxDocument
from openpyxl import Workbook

from app.cleaners import clean_text
from app.deduplication import (
    DuplicateCandidate,
    choose_authoritative,
    content_hash,
    hamming_distance,
    is_near_duplicate,
    simhash,
)
from app.parsers.docx import DocxParser
from app.parsers.html import HtmlParser
from app.parsers.pdf import PdfParser
from app.parsers.registry import ParserRegistry
from app.parsers.text import TextParser
from app.parsers.xlsx import XlsxParser
from app.parsers.zip_metadata import ZipMetadataParser


def test_html_parser_removes_boilerplate_and_preserves_structure() -> None:
    parsed = HtmlParser().parse(
        b"""
        <html><body><nav>navigation</nav><article>
        <h1>Policy title</h1><p>First paragraph.</p>
        <ul><li>Condition A</li></ul>
        <table><tr><th>Item</th><th>Value</th></tr><tr><td>Funding</td><td>100</td></tr></table>
        <a href="/attachment.pdf">Attachment</a>
        </article><footer>footer</footer></body></html>
        """
    )
    assert "navigation" not in parsed.text
    assert "# Policy title" in parsed.text
    assert "- Condition A" in parsed.text
    assert parsed.tables[0].headers == ("Item", "Value")
    assert parsed.metadata["links"][0]["href"] == "/attachment.pdf"


def test_pdf_parser_removes_repeated_header_and_footer() -> None:
    pymupdf = import_module("pymupdf")
    document = pymupdf.open()
    for body in ("First page body", "Second page body"):
        page = document.new_page()
        page.insert_text((72, 72), "Repeated header")
        page.insert_text((72, 100), body)
        page.insert_text((72, 740), "Repeated footer")
    content = document.tobytes()
    document.close()
    parsed = PdfParser().parse(content)
    assert "Repeated header" not in parsed.text
    assert "Repeated footer" not in parsed.text
    assert "First page body" in parsed.text
    assert parsed.metadata["page_count"] == 2


def test_pdf_parser_flags_image_only_document_for_ocr() -> None:
    pymupdf = import_module("pymupdf")
    document = pymupdf.open()
    document.new_page()
    content = document.tobytes()
    document.close()
    assert PdfParser().parse(content).requires_ocr


def test_docx_parser_extracts_heading_paragraph_and_table() -> None:
    document = DocxDocument()
    document.add_heading("Application conditions", level=1)
    document.add_paragraph("Registered enterprises may apply.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Field"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Region"
    table.rows[1].cells[1].text = "Sichuan"
    buffer = BytesIO()
    document.save(buffer)
    parsed = DocxParser().parse(buffer.getvalue())
    assert parsed.sections[0].title == "Application conditions"
    assert "Registered enterprises" in parsed.text
    assert parsed.tables[0].rows[0] == ("Region", "Sichuan")


def test_xlsx_parser_emits_sheet_markdown_and_json_rows() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Indicators"
    sheet.append(["name", "value"])
    sheet.append(["projects", 12])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    parsed = XlsxParser().parse(buffer.getvalue())
    assert "# Indicators" in parsed.text
    assert parsed.metadata["sheets"][0]["rows"][0] == {"name": "projects", "value": "12"}


def test_text_and_zip_metadata_parsers() -> None:
    assert TextParser().parse("政策正文".encode("gb18030")).text == "政策正文"
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("safe/readme.txt", "body")
        archive.writestr("../unsafe.txt", "body")
    parsed = ZipMetadataParser().parse(buffer.getvalue())
    assert parsed.metadata["entries"][0]["unsafe_path"] is False
    assert parsed.metadata["entries"][1]["unsafe_path"] is True


def test_parser_registry_rejects_unknown_extension() -> None:
    try:
        ParserRegistry().get("document.exe")
    except ValueError as exc:
        assert "unsupported parser extension" in str(exc)
    else:
        raise AssertionError("unknown extension must fail")


def test_cleaning_hashing_simhash_and_authoritative_selection() -> None:
    cleaned = clean_text("正文\n\n正文\n\n打印本页")
    assert cleaned == "正文"
    assert content_hash("a b") == content_hash("ab")
    left = simhash("software policy support funding")
    right = simhash("software policy support funding")
    assert hamming_distance(left, right) == 0
    assert is_near_duplicate("政策支持软件企业", "政策支持软件企业")
    selected = choose_authoritative(
        [
            DuplicateCandidate("local", False, 100, date(2026, 1, 2), 500),
            DuplicateCandidate("official", True, 10, date(2026, 1, 1), 300),
        ]
    )
    assert selected.document_id == "official"
